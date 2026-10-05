"""Traceable table intake and biological-unit preparation; no silent imputation."""
from __future__ import annotations
import csv
from pathlib import Path
import numpy as np
import pandas as pd
from .common import digest


def read_table(path,sheet=None,header_row=1):
    path=Path(path).resolve(); header_row=int(header_row)
    if header_row<1: raise ValueError('header_row is 1-based and must be positive')
    if path.suffix.lower()=='.xlsx':
        import openpyxl
        book=openpyxl.load_workbook(path,data_only=True,read_only=True)
        if sheet is None:
            if len(book.sheetnames)!=1: raise ValueError('Select a worksheet explicitly: '+', '.join(book.sheetnames))
            sheet=book.sheetnames[0]
        if sheet not in book.sheetnames: raise ValueError('Unknown worksheet '+str(sheet))
        rows=list(book[sheet].iter_rows(values_only=True)); book.close()
        if len(rows)<header_row: raise ValueError('Header row beyond data')
        header=list(rows[header_row-1]); data=rows[header_row:]; loc=list(range(header_row+1,len(rows)+1))
    elif path.suffix.lower() in ['.csv','.tsv','.txt']:
        rows=[]; loc=[]
        with path.open(encoding='utf-8-sig',newline='') as f:
            reader=csv.reader(f,delimiter='\t' if path.suffix.lower() in ['.tsv','.txt'] else ',')
            previous=0
            for row in reader:
                rows.append(row); loc.append(previous+1); previous=reader.line_num
        if len(rows)<header_row: raise ValueError('Header row beyond data')
        header=rows[header_row-1]; data=rows[header_row:]; loc=loc[header_row:]; sheet=None
    else: raise ValueError('Supported raw tables: CSV, TSV, XLSX; convert legacy XLS or instrument-native files explicitly')
    header=[str(v).strip() if v is not None else '' for v in header]
    # Ignore entirely empty trailing spreadsheet columns, never populated unnamed columns.
    while header and not header[-1] and all(len(r)<len(header) or r[len(header)-1] in [None,''] for r in data):
        header.pop(); data=[r[:len(header)] for r in data]
    if not header or any(not x for x in header): raise ValueError('Blank column header')
    if len(set(header))!=len(header): raise ValueError('Duplicate headers; resolve source ambiguity before analysis')
    if any(x.startswith('__') for x in header): raise ValueError('Column names beginning __ are reserved for source lineage')
    if any(len(r)!=len(header) for r in data): raise ValueError('Ragged table rows do not match header width')
    df=pd.DataFrame(data,columns=header).replace('',np.nan)
    df['__source_row']=loc; df=df.dropna(how='all',subset=header)
    source={'path':str(path),'sha256':digest(path.read_bytes()),'sheet':sheet,'header_row':header_row,'rows':len(df),'columns':header}
    return df,source


def profile(path,sheet=None,header_row=1):
    if Path(path).suffix.lower()=='.xlsx' and sheet is None:
        import openpyxl
        with_path=openpyxl.load_workbook(path,read_only=True,data_only=True); names=with_path.sheetnames; with_path.close()
        return {'file':str(Path(path).resolve()),'sheets':[profile(path,x,header_row) for x in names],'status':'profile_only; no experimental design inferred'}
    df,src=read_table(path,sheet,header_row)
    return {'source':src,'columns':{c:{'missing':int(df[c].isna().sum()),'unique':int(df[c].nunique()),'numeric_values':int(pd.to_numeric(df[c],errors='coerce').notna().sum())} for c in src['columns']},'duplicate_rows':int(df.duplicated(subset=src['columns']).sum()),'status':'profile_only; raw values not sent to remote services'}


def load_experiment(config,base):
    inp=config['input']; path=Path(base)/inp['data']; df,source=read_table(path,inp.get('sheet'),inp.get('header_row',1)); sources=[source]
    mapping=config.get('columns',{}); df=df.rename(columns={v:k for k,v in mapping.items() if v in df.columns})
    meta=None
    if inp.get('samples'):
        meta,sm=read_table(Path(base)/inp['samples'],inp.get('sample_sheet'),inp.get('sample_header_row',1)); sources.append(sm)
        meta=meta.rename(columns={v:k for k,v in config.get('sample_columns',mapping).items() if v in meta.columns})
        if 'sample' not in meta or meta['sample'].isna().any() or meta['sample'].duplicated().any(): raise ValueError('Sample metadata must have one nonempty row per sample')
        meta['sample']=meta['sample'].astype(str).str.strip()
        if meta['sample'].duplicated().any(): raise ValueError('Sample IDs collide after trimming')
    if inp.get('layout','long')=='wide':
        if meta is None: raise ValueError('Wide matrices require sample metadata; numeric columns are not guessed as samples')
        sample_names=meta['sample'].tolist()
        if any(x not in df for x in sample_names): raise ValueError('Metadata sample missing from matrix: '+str([x for x in sample_names if x not in df]))
        df=df.melt(id_vars=[c for c in df if c not in sample_names],value_vars=sample_names,var_name='sample',value_name='value')
        df['__source_column']=df['sample']
    else: df['__source_column']=mapping.get('value','value')
    if 'sample' not in df: raise ValueError('Missing sample column')
    if df['sample'].isna().any(): raise ValueError('Empty sample ID')
    df['sample']=df['sample'].astype(str).str.strip()
    if meta is not None:
        absent=sorted(set(meta['sample'])-set(df['sample']))
        if absent: raise ValueError('Sample manifest includes samples with no measurement rows: '+', '.join(absent[:20])+'; document missing units or an explicit exclusion in a new input version before inference')
        overlap=[c for c in meta if c in df and c not in ['sample','__source_row']]
        if overlap: raise ValueError('Conflicting data/metadata columns: '+', '.join(overlap)+'; keep sample attributes in one table')
        meta=meta.rename(columns={'__source_row':'__metadata_row'})
        df=df.merge(meta,on='sample',how='left',validate='many_to_one',indicator=True)
        if (df['_merge']!='both').any(): raise ValueError('Some data samples have no metadata')
        df=df.drop(columns='_merge')
    if 'unit' not in df and config.get('study',{}).get('sample_is_biological_unit') is True: df['unit']=df['sample']
    if df.empty: raise ValueError('No observations in input')
    required=['feature','sample','group','unit','value']
    if any(c not in df for c in required): raise ValueError('Required columns missing: '+', '.join(c for c in required if c not in df))
    for col in ['feature','sample','group','unit']:
        if df[col].isna().any(): raise ValueError('Empty identifier in '+col)
        df[col]=df[col].astype(str).str.strip()
        if (df[col]=='').any(): raise ValueError('Blank identifier in '+col)
    for col in ['pair','batch']:
        if col in df:
            if df[col].isna().any(): raise ValueError('Missing declared '+col+' values')
            df[col]=df[col].astype(str).str.strip()
    raw=df['value'].copy(); sentinels=config.get('analysis',{}).get('missing_values',[])
    df.loc[df['value'].isin(sentinels),'value']=np.nan
    numeric=pd.to_numeric(df['value'],errors='coerce')
    invalid=df['value'].notna() & numeric.isna()
    if invalid.any(): raise ValueError('Non-numeric value cells require an explicit missing-value rule; source rows: '+str(df.loc[invalid,'__source_row'].tolist()[:12]))
    if np.isinf(numeric).any(): raise ValueError('Infinite values in input')
    df['value']=numeric.astype(float); df['__input_value']=raw
    if 'time' in df:
        df['time']=pd.to_numeric(df['time'],errors='raise')
        if df['time'].isna().any() or not np.isfinite(df['time']).all(): raise ValueError('time must contain finite numeric values with declared units')
    unique_key=['sample','feature']+([ 'technical'] if 'technical' in df else [])+(['time'] if 'time' in df else [])
    if df.duplicated(unique_key).any(): raise ValueError('Duplicate measurement key; supply technical replicate IDs or resolve duplicate import')
    df['__source_file']=source['path']; df['__source_sheet']=source['sheet'] or ''
    return df,sources


def prepare_units(df,config):
    study=config['study']; analysis=config.get('analysis',{}); issues=[]
    if study.get('data_scale')=='raw_counts':
        values=df['value']
        invalid=values.isna() | (values<0) | (values!=np.floor(values))
        if invalid.any(): issues.append({'severity':'blocking','code':'invalid_raw_count_cells','detail':'Raw counts must be nonnegative integers in every supplied cell before technical aggregation; missing or invalid source rows: '+str(df.loc[invalid,'__source_row'].tolist()[:20])})
    if study.get('independent_units_confirmed') is not True: issues.append({'severity':'blocking','code':'unit_definition','detail':'Declare the true experimental unit before inferential analysis'})
    design=study.get('design')
    if 'time' in df and not study.get('time_unit'): raise ValueError('Declare study.time_unit for a numeric time column')
    for col in ['symbol','accession']:
        if col in df and (df.groupby('feature')[col].nunique()>1).any():
            issues.append({'severity':'blocking','code':'feature_identity_conflict','detail':'One feature maps to multiple '+col+' values; resolve identifiers before inference'})
    if design not in ['independent','paired','longitudinal']: raise ValueError('design must be independent, paired or longitudinal')
    if design in ['independent','longitudinal'] and (df.groupby('unit')['group'].nunique()>1).any(): issues.append({'severity':'blocking','code':'unit_cross_group','detail':'Independent-unit IDs occur in multiple groups; do not infer pairing or independence'})
    keys=['feature','group','unit']+(['time'] if 'time' in df else [])
    sizes=df.groupby(keys,dropna=False).size(); aggregation=analysis.get('technical_aggregation','none')
    if sizes.max()>1 and aggregation not in ['mean','median','sum']:
        issues.append({'severity':'blocking','code':'technical_replicates','detail':'Multiple measurements per biological unit require explicit technical aggregation'})
        aggregation='mean'  # QC summary only; inference is disabled below.
    if aggregation=='sum' and study.get('data_scale')!='raw_counts': raise ValueError('Technical summation is only supported for declared count libraries')
    for col in ['batch','pair','symbol','accession']:
        if col in df and (df.groupby(keys)[col].nunique()>1).any(): issues.append({'severity':'blocking','code':'conflicting_'+col,'detail':'An experimental unit has conflicting '+col+' annotations'})
    operations={'value':aggregation if aggregation in ['mean','median','sum'] else 'mean','sample':lambda x:list(dict.fromkeys(x)), '__source_row':lambda x:sorted(set(int(z) for z in x)), '__source_column':lambda x:list(dict.fromkeys(str(z) for z in x))}
    operations.update({c:'first' for c in ['batch','pair','symbol','accession'] if c in df})
    units=df.groupby(keys,dropna=False,sort=False).agg(operations).reset_index()
    units['observed_technical_n']=df.groupby(keys,dropna=False,sort=False)['value'].count().to_numpy()
    if aggregation=='sum': units.loc[units['observed_technical_n']==0,'value']=np.nan
    units['input_value']=units['value']
    scale=study.get('data_scale')
    if scale not in ['linear','log2','ct','raw_counts']: raise ValueError('Declare data_scale: linear, log2, ct or raw_counts')
    if scale=='ct':
        qpcr=config.get('qpcr',{}); refs=qpcr.get('reference_features',[])
        if not refs or qpcr.get('equal_efficiency_confirmed') is not True: raise ValueError('Ct analysis needs reference_features and equal_efficiency_confirmed; otherwise use an explicit efficiency model')
        units=normalize_reference(units,refs,mode='ct'); scale='log2'
    norm=config.get('normalization',{})
    if norm.get('method')=='ratio':
        if scale!='linear': raise ValueError('Ratio normalization requires linear background-corrected measurements')
        units=normalize_reference(units,norm.get('reference_features',[]),mode='ratio')
    transform=analysis.get('transform','none')
    if transform=='log2':
        if scale!='linear': raise ValueError('log2 transform is only allowed on declared linear values')
        if (units['value'].dropna()<=0).any(): raise ValueError('Nonpositive values cannot be logged; no automatic pseudocount or zero imputation')
        units['value']=np.log2(units['value']); scale='log2'
    elif transform!='none': raise ValueError('Unsupported transform')
    # Detect hidden imputation floor without deleting or automatically imputing anything.
    values=units['value'].dropna()
    if len(values)>=20:
        floor=values.min(); floor_n=int((values==floor).sum())
        if floor_n>=max(4,len(values)*.15): issues.append({'severity':'warning','code':'repeated_floor','detail':f'{floor_n} values equal the minimum {floor}; check censoring/imputation before mechanistic use'})
    if 'batch' in units and units['batch'].nunique()>1:
        unit_meta=units[['unit','group','batch']].drop_duplicates()
        if (unit_meta.groupby('batch')['group'].nunique()==1).all(): issues.append({'severity':'blocking','code':'batch_group_confounding','detail':'Every batch contains only one group; group and batch effects are not separable'})
        elif not analysis.get('adjust_batch',False): issues.append({'severity':'blocking','code':'batch_unadjusted','detail':'Multiple batches present; declare adjust_batch or provide a documented alternative model'})
    return units,issues,scale


def normalize_reference(units,refs,mode):
    if not refs: raise ValueError('Reference features required')
    keys=['group','unit']+(['time'] if 'time' in units else [])
    reference=units[units['feature'].isin(refs)].groupby(keys,dropna=False)['value'].agg(['mean','count']).reset_index()
    reference=reference.rename(columns={'mean':'reference_value','count':'reference_n'})
    targets=units[~units['feature'].isin(refs)].merge(reference,on=keys,how='left',validate='many_to_one')
    if (targets['reference_n'].fillna(0)!=len(refs)).any() or targets['reference_value'].isna().any(): raise ValueError('Missing reference measurement for some biological units')
    if mode=='ct': targets['delta_ct']=targets['value']-targets['reference_value']; targets['value']=-targets['delta_ct']
    else:
        if (targets['reference_value']<=0).any(): raise ValueError('Reference intensity must be positive')
        targets['value']=targets['value']/targets['reference_value']
    return targets
